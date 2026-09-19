"""Problem Statement 3 predict routes - and the task-runner core the CLI shares with them.

Two halves, in this order:

1. **Core** (no FastAPI in the call path): resolving a registered ``Task``, loading its committed
   artefact, running one input file, collecting the inputs under a directory, inferring the task
   from an input, and rendering the organiser CSV. ``scripts/ps3_predict.py``, the root
   ``predict.py`` and ``scripts/ps3_submission.py`` import exactly these, so the bytes the app
   offers for download and the bytes the CLI writes come out of the same function.
2. **Router** (`/api/ps3/...`), mounted by :func:`nebulax.api.main.create_app`.

Upload sessions
---------------
68 rail files are ~1.1 GB, so a browser never sends them in one request. ``POST
/api/ps3/{task}/predict`` takes a bounded batch, processes the files **sequentially** (one model
in memory, one file on disk at a time), appends the rows to a session and hands back a session
token; the client posts the next batch with the same token. ``GET /api/ps3/results/{token}.csv``
renders everything accumulated so far and validates it with
:func:`nebulax.ps3.submission.validate_csv` before it leaves the process.

Nothing here fits a model or knows what a subsystem *is*: the four task modules
(``nebulax/ps3/{door,acv,rail,shm}.py``) own that, behind the protocol in ``docs/ps3_contract.md``.
A task whose module or artefact is missing answers **503**, never a traceback, so the app runs
while the model agents are still writing.
"""

from __future__ import annotations

import csv
import io
import json
import math
import os
import re
import shutil
import tempfile
import time
import uuid
from dataclasses import dataclass, field
from pathlib import Path
from typing import Any, Iterable, Mapping, Sequence

import numpy as np
from fastapi import APIRouter, Body, File, Form, HTTPException, Request, Response, UploadFile
from fastapi.concurrency import run_in_threadpool
from starlette.datastructures import UploadFile as _StarletteUpload

from nebulax.ps3 import common
from nebulax.ps3.common import (
    ACCEPTED_SUFFIXES,
    OUTPUT_FILENAMES,
    TASK_LABELS,
    TASK_NAMES,
    PredictionResult,
    acv_car_ids,
    get_task,
    natural_key,
)
from nebulax.ps3.submission import CSV_HEADERS, validate_csv

__all__ = [
    # core
    "TaskUnavailable",
    "MAX_UPLOAD_BYTES",
    "MAX_BATCH_FILES",
    "SESSION_TTL_S",
    "resolve_task",
    "task_model",
    "run_file",
    "predict_paths",
    "collect_inputs",
    "infer_task",
    "render_csv",
    "write_csv",
    "cv_path",
    "cv_summary",
    # router
    "router",
    "reset_sessions",
]

#: Hard per-file upload cap (the biggest released rail file is ~17 MB).
MAX_UPLOAD_BYTES: int = 64 * 1024 * 1024

#: How long an idle upload session (and its temp files) survives.
SESSION_TTL_S: float = 2 * 60 * 60.0

#: Files one request may carry. The 68 rail files are ~1.1 GB, so the page uploads them in
#: batches; ``GET /api/ps3/tasks`` publishes this number so the client can size its batches.
MAX_BATCH_FILES: int = 32

#: Token shape: what we mint, and the only thing we will look up.
_TOKEN_RE = re.compile(r"^[A-Za-z0-9_-]{8,64}$")

#: Door's optional trailing column (`Door Info Kit`; the judge ignores it).
_DOOR_EXTRA = "confidence"


class TaskUnavailable(RuntimeError):
    """A PS3 task cannot run yet: its module is missing, or its artefact has not been fitted."""


# ======================================================================================
# Core - shared with the CLI (no FastAPI below this line until the router section)
# ======================================================================================


def _check_task(name: str) -> str:
    key = str(name).strip().lower()
    if key not in TASK_NAMES:
        raise KeyError(f"unknown PS3 task {name!r}; expected one of {list(TASK_NAMES)}")
    return key


def _model_dir() -> Path | None:
    env = os.environ.get("NEBULAX_PS3_MODELS", "").strip()
    return Path(env).expanduser() if env else None


def _results_dir() -> Path:
    env = os.environ.get("NEBULAX_PS3_RESULTS", "").strip()
    return Path(env).expanduser() if env else common.RESULTS_DIR


def resolve_task(name: str) -> Any:
    """The registered task object, or :class:`TaskUnavailable` with a readable message.

    ``get_task`` imports ``nebulax.ps3.<name>`` lazily; any failure of that import (the module is
    not written yet, or it blows up) becomes one ``TaskUnavailable`` the caller turns into a 503.
    """
    key = _check_task(name)
    try:
        return get_task(key)
    except KeyError as exc:
        raise TaskUnavailable(
            f"PS3 task {key!r} is not available yet: {exc}. "
            f"Fit it with `python scripts/ps3_train.py --task {key}`."
        ) from exc
    except Exception as exc:  # a half-written task module must not 500 the whole app
        raise TaskUnavailable(f"PS3 task {key!r} could not be imported: {type(exc).__name__}: {exc}") from exc


#: ``{(task, path, mtime_ns): model}`` - one artefact is reused across a whole batch/directory.
_MODEL_CACHE: dict[tuple[str, str, int], Any] = {}


def task_model(name: str) -> Any | None:
    """The committed artefact for a task, cached by file mtime; ``None`` when there is none.

    ``None`` is not an error: a task may be a pure rule (the ACV ranker is) and
    ``Task.predict(feats, None)`` is contractually "load whatever you need yourself".
    """
    key = _check_task(name)
    pkl = common.model_path(key, model_dir=_model_dir())
    try:
        stat = pkl.stat()
    except OSError:
        return None
    ck = (key, str(pkl), int(stat.st_mtime_ns))
    if ck not in _MODEL_CACHE:
        _MODEL_CACHE.clear()  # one artefact at a time: these are up to 5 MB each
        _MODEL_CACHE[ck] = common.load_model(key, model_dir=_model_dir())
    return _MODEL_CACHE[ck]


def run_file(task: Any, path: Path | str, model: Any = None) -> PredictionResult:
    """``load`` -> ``featurise`` -> ``predict`` for one input, with ``file_id`` defaulted.

    Same three calls as :meth:`nebulax.ps3.common.BaseTask.run`, spelled out so it also works for
    a task that implements the protocol without inheriting from ``BaseTask``.
    """
    p = Path(path)
    result = task.predict(task.featurise(task.load(p)), model)
    if not isinstance(result, PredictionResult):
        raise TypeError(f"{task.name}.predict returned {type(result).__name__}, expected PredictionResult")
    if not result.file_id and task.name != "door":
        result.file_id = p.name
    return result


def predict_paths(
    name: str,
    paths: Sequence[Path | str],
    *,
    model: Any = None,
    on_file: Any = None,
) -> tuple[list[dict[str, Any]], list[dict[str, Any]], list[dict[str, str]]]:
    """Run a task over several inputs **sequentially**; the CLI's whole body.

    Returns ``(rows, explanations, errors)``: organiser rows in input order, one
    ``Explanation.as_dict()`` per successfully predicted file, and ``{"file", "message"}`` for
    every input that failed. A missing artefact raises :class:`TaskUnavailable` (it would fail
    identically for every file, so it is not a per-file error).
    """
    task = resolve_task(name)
    if model is None:
        model = task_model(name)
    rows: list[dict[str, Any]] = []
    explanations: list[dict[str, Any]] = []
    errors: list[dict[str, str]] = []
    for path in paths:
        p = Path(path)
        try:
            result = run_file(task, p, model)
            rows.extend(task.to_rows(result))
            explanations.append(task.explain(result).as_dict())
        except FileNotFoundError as exc:
            if p.exists():  # the input is there; something the task needed is not
                raise TaskUnavailable(f"PS3 task {name!r} cannot predict yet: {exc}") from exc
            errors.append({"file": p.name, "message": f"cannot read {p.name}: {exc}"})
        except Exception as exc:
            errors.append({"file": p.name, "message": f"{type(exc).__name__}: {exc}"})
        else:
            if on_file is not None:
                on_file(p, result)
    return rows, explanations, errors


def collect_inputs(name: str, target: Path | str) -> list[Path]:
    """The input files for a task under ``target`` (a file, or a directory), in natural order."""
    key = _check_task(name)
    p = Path(target)
    suffixes = ACCEPTED_SUFFIXES[key]
    if p.is_file():
        return [p]
    if not p.is_dir():
        raise FileNotFoundError(f"no such input: {p}")
    files = [
        f
        for f in p.iterdir()
        if f.is_file() and f.suffix.lower() in suffixes and not f.name.startswith("~$") and not f.name.startswith(".")
    ]
    if not files:
        raise FileNotFoundError(f"{p} holds no {' or '.join(suffixes)} file for task {key!r}")
    return sorted(files, key=lambda f: natural_key(f.name))


# --------------------------------------------------------------------------------------
# Task inference (root predict.py: the info kits' `--input/--output` interface)
# --------------------------------------------------------------------------------------

#: First two columns of the door stream (`Door/Train.csv`, `Door/Test.csv`).
_DOOR_HEAD = ("datetime", "motor current")


def _sniff_csv(path: Path) -> str:
    """``door`` / ``rail`` / ``shm`` from one csv's first lines, by shape not by name."""
    with open(path, "r", encoding="utf-8-sig", errors="replace", newline="") as fh:
        first = fh.readline()
        second = fh.readline()
    header = [c.strip().lower() for c in next(csv.reader([first]), [])]
    if len(header) >= 2 and header[0].startswith(_DOOR_HEAD[0]) and header[1].startswith(_DOOR_HEAD[1]):
        return "door"
    n_first = len(next(csv.reader([first]), []))
    n_second = len(next(csv.reader([second]), [])) if second.strip() else n_first
    if max(n_first, n_second) > 8:
        return "rail"  # the released rail files are 129 columns (speed + 64 boxes x 2)
    try:
        float(first.strip().split(",")[0])
    except ValueError as exc:
        raise ValueError(
            f"cannot tell which PS3 subsystem {path.name} belongs to (header {header[:4]}); pass --task"
        ) from exc
    return "shm"  # headerless single-column dynamic stress


def infer_task(target: Path | str) -> str:
    """Guess the subsystem from an input file or directory (``--task`` always wins over this).

    xlsx/xls -> ``acv``; the door header -> ``door``; a wide csv -> ``rail``; a headerless
    single-column csv -> ``shm``. Raises ``ValueError`` with a readable message when it cannot
    tell, so the caller can say "pass --task".
    """
    p = Path(target)
    if p.is_file():
        if p.suffix.lower() in (".xlsx", ".xls"):
            return "acv"
        if p.suffix.lower() != ".csv":
            raise ValueError(f"{p.name}: PS3 inputs are .csv (door/rail/shm) or .xlsx (acv)")
        return _sniff_csv(p)
    if not p.is_dir():
        raise FileNotFoundError(f"no such input: {p}")
    xlsx = [f for f in p.iterdir() if f.is_file() and f.suffix.lower() in (".xlsx", ".xls") and not f.name.startswith("~$")]
    csvs = sorted(
        (f for f in p.iterdir() if f.is_file() and f.suffix.lower() == ".csv" and not f.name.startswith(".")),
        key=lambda f: natural_key(f.name),
    )
    if xlsx and not csvs:
        return "acv"
    if not csvs:
        raise FileNotFoundError(f"{p} holds no PS3 input file (.csv or .xlsx)")
    votes = {_sniff_csv(f) for f in csvs[:3]}
    if len(votes) != 1:
        raise ValueError(f"{p} mixes PS3 subsystems ({sorted(votes)}); pass --task")
    return votes.pop()


# --------------------------------------------------------------------------------------
# CSV rendering - the single definition of "the submission bytes"
# --------------------------------------------------------------------------------------


def _cell(value: Any) -> str:
    """One CSV cell. Floats use ``repr`` (shortest round-trip), so the CLI and the API agree."""
    if isinstance(value, (np.generic,)):
        value = value.item()
    if isinstance(value, bool):
        return str(value)
    if isinstance(value, float):
        return repr(value)
    return "" if value is None else str(value)


def render_csv(name: str, rows: Iterable[dict[str, Any]]) -> str:
    """The organiser CSV for ``rows`` as text (LF line endings, no BOM, header order frozen).

    Only the contract's columns are written, in order, plus door's optional ``confidence`` when
    **every** row carries one. Missing keys become empty cells - the validator will say so rather
    than a silently wrong submission going out.
    """
    key = _check_task(name)
    materialised = [dict(r) for r in rows]
    header = list(CSV_HEADERS[key])
    if key == "door" and materialised and all(_DOOR_EXTRA in r for r in materialised):
        header.append(_DOOR_EXTRA)
    buf = io.StringIO(newline="")
    writer = csv.writer(buf, lineterminator="\n")
    writer.writerow(header)
    for row in materialised:
        writer.writerow([_cell(row.get(col)) for col in header])
    return buf.getvalue()


def write_csv(name: str, rows: Iterable[dict[str, Any]], out_path: Path | str) -> Path:
    """:func:`render_csv` to a file (utf-8, LF). Byte-identical to what the API serves."""
    out = Path(out_path)
    out.parent.mkdir(parents=True, exist_ok=True)
    with open(out, "w", encoding="utf-8", newline="") as fh:
        fh.write(render_csv(name, rows))
    return out


# --------------------------------------------------------------------------------------
# CV results
# --------------------------------------------------------------------------------------


def cv_path(name: str) -> Path:
    """``results/ps3/<task>_cv.json`` (``NEBULAX_PS3_RESULTS`` overrides the directory)."""
    return _results_dir() / f"{_check_task(name)}_cv.json"


#: Scalar top-level entries worth showing next to a task in the app's subsystem picker.
_CV_SCALAR_KEYS = (
    "task",
    "scheme",
    "metric",
    "score",
    "n_folds",
    "n_blocks",
    "n_files",
    "n_cases",
    "seeds",
    "model",
    "feature_version",
    "git_rev",
    "generated_at",
    "saved_at",
)

#: Objects whose scalar entries are kept, one level deep (their per-fold arrays are dropped).
_CV_NESTED_KEYS = ("headline", "summary", "winner")


def _scalars(obj: Mapping[str, Any], limit: int = 16) -> dict[str, Any]:
    """The scalar (and short flat list) entries of a mapping - never a per-fold array."""
    out: dict[str, Any] = {}
    for key, value in obj.items():
        keep = isinstance(value, (str, int, float, bool)) or (
            isinstance(value, list)
            and len(value) <= 8
            and all(isinstance(x, (str, int, float, bool)) for x in value)
        )
        if keep:
            out[str(key)] = value
        if len(out) >= limit:
            break
    return out


def cv_summary(name: str) -> dict[str, Any] | None:
    """A small summary of ``<task>_cv.json`` for the task list (``None`` when it is not written).

    Keeps the scheme/metric/seeds/git-rev line plus the scalar entries of ``headline``,
    ``summary`` and ``winner``, so a five-fold array never travels to the browser on a page load.
    ``GET /api/ps3/{task}/cv`` serves the whole file for anyone who wants the folds.
    """
    path = cv_path(name)
    if not path.exists():
        return None
    try:
        data = json.loads(path.read_text(encoding="utf-8"))
    except (OSError, ValueError):
        return None
    if not isinstance(data, dict):
        return None
    out = _scalars({k: v for k, v in data.items() if k in _CV_SCALAR_KEYS})
    for key in _CV_NESTED_KEYS:
        value = data.get(key)
        if isinstance(value, dict):
            sub = _scalars(value)
            if sub:
                out[key] = sub

    # SHM's ``*_cv.json`` is deliberately the frozen stats-only baseline.  Its honest selected
    # estimate lives under ``*_ladder.json:nested``; use that for the tab instead of displaying
    # "no headline" (the full baseline CV remains available at the dedicated /cv route).
    if name == "shm" and "headline" not in out and "score" not in out:
        ladder = _results_dir() / "shm_ladder.json"
        try:
            nested = json.loads(ladder.read_text(encoding="utf-8")).get("nested", {})
        except (OSError, ValueError, AttributeError):
            nested = {}
        if isinstance(nested, dict) and isinstance(nested.get("score"), (int, float)):
            out["metric"] = "mape_score"
            out["headline"] = _scalars(nested)
    return out or None


# ======================================================================================
# Router
# ======================================================================================


def _upload_root() -> Path:
    env = os.environ.get("NEBULAX_PS3_UPLOADS", "").strip()
    return Path(env).expanduser() if env else common.CACHE_DIR / "uploads"


@dataclass
class _Session:
    """One accumulating upload session: the rows so far plus the temp files they came from."""

    token: str
    task: str
    created: float
    touched: float
    rows: list[dict[str, Any]] = field(default_factory=list)
    files: list[str] = field(default_factory=list)
    expected_cars: dict[str, list[str]] = field(default_factory=dict)
    n_batches: int = 0

    @property
    def dir(self) -> Path:
        return _upload_root() / self.token


#: Live sessions, keyed by token. Module level because the router is a module-level object.
_SESSIONS: dict[str, _Session] = {}


def reset_sessions() -> None:
    """Drop every session and its temp files (tests; and a clean process restart)."""
    for sess in list(_SESSIONS.values()):
        shutil.rmtree(sess.dir, ignore_errors=True)
    _SESSIONS.clear()


#: Last time :func:`_sweep` looked for orphan directories (a restart loses the session dict).
_LAST_ORPHAN_SWEEP: float = 0.0


def _sweep(now: float | None = None) -> None:
    """Expire idle sessions and delete their temp files (including orphans from a past process)."""
    global _LAST_ORPHAN_SWEEP
    t = time.time() if now is None else now
    for token, sess in list(_SESSIONS.items()):
        if t - sess.touched > SESSION_TTL_S:
            shutil.rmtree(sess.dir, ignore_errors=True)
            _SESSIONS.pop(token, None)
    if t - _LAST_ORPHAN_SWEEP < 60.0:
        return
    _LAST_ORPHAN_SWEEP = t
    root = _upload_root()
    try:
        entries = list(root.iterdir()) if root.is_dir() else []
    except OSError:  # pragma: no cover - a vanished cache dir is not an API failure
        return
    for entry in entries:
        if entry.name in _SESSIONS or not entry.is_dir():
            continue
        try:
            stale = t - entry.stat().st_mtime > SESSION_TTL_S
        except OSError:  # pragma: no cover
            continue
        if stale:
            shutil.rmtree(entry, ignore_errors=True)


def _task_or_404(name: str) -> str:
    try:
        return _check_task(name)
    except KeyError as exc:
        raise HTTPException(status_code=404, detail=str(exc)) from exc


def _session_or_404(token: str) -> _Session:
    if not _TOKEN_RE.match(str(token or "")):
        raise HTTPException(status_code=404, detail=f"unknown session {token!r}")
    sess = _SESSIONS.get(token)
    if sess is None:
        raise HTTPException(status_code=404, detail=f"unknown or expired session {token!r}")
    return sess


async def _save_upload(upload: UploadFile, dest: Path) -> int:
    """Stream one upload to ``dest`` in 1 MB chunks, refusing anything over the cap."""
    total = 0
    try:
        with open(dest, "wb") as fh:
            while True:
                chunk = await upload.read(1 << 20)
                if not chunk:
                    break
                total += len(chunk)
                if total > MAX_UPLOAD_BYTES:
                    raise HTTPException(
                        status_code=413,
                        detail=(
                            f"{dest.name} is larger than the {MAX_UPLOAD_BYTES // (1024 * 1024)} MB "
                            "per-file limit; upload the files in smaller batches"
                        ),
                    )
                fh.write(chunk)
    except HTTPException:
        dest.unlink(missing_ok=True)
        raise
    except OSError as exc:
        dest.unlink(missing_ok=True)
        raise HTTPException(status_code=400, detail=f"could not store {dest.name}: {exc}") from exc
    finally:
        await upload.close()
    return total


router = APIRouter(prefix="/api/ps3", tags=["ps3"])


def inspect_acv_workbook(path: Path) -> dict[str, Any]:
    """Read only the first worksheet's header and a few rows before an ACV file is queued.

    This is a format check, not a prediction or a guarantee that every later value is usable.
    The full loader remains authoritative when RUN is pressed.
    """
    import pandas as pd
    from openpyxl import load_workbook

    from nebulax.ps3.acv_features import CAR_COL_RE, SYNONYMS

    try:
        book = load_workbook(path, read_only=True, data_only=True)
    except Exception as exc:
        return {"valid": False, "errors": [f"not a readable .xlsx workbook: {exc}"], "summary": {}}
    try:
        sheet = book.worksheets[0] if book.worksheets else None
        if sheet is None:
            return {"valid": False, "errors": ["workbook has no worksheet"], "summary": {}}
        rows = list(sheet.iter_rows(min_row=1, max_row=8, values_only=True))
        headers = [str(v).strip() if v is not None else "" for v in (rows[0] if rows else ())]
        errors: list[str] = []
        time_idx = next((i for i, name in enumerate(headers) if name.lower() == "time"), None)
        if time_idx is None:
            errors.append("missing Time column")
        cars: set[str] = set()
        params: set[str] = set()
        for name in headers:
            match = CAR_COL_RE.match(name)
            if match:
                cars.add(match.group(1))
                params.add(match.group(2).strip())
        if len(cars) != 8:
            errors.append(f"expected headers for 8 cars; found {len(cars)}")
        known = {syn for synonyms in SYNONYMS.values() for syn in synonyms}
        recognized = sorted(params & known)
        if not recognized:
            errors.append("no recognized ACV sensor fields after the 'Car NN - ' prefixes")
        if time_idx is not None:
            samples = [r[time_idx] for r in rows[1:] if len(r) > time_idx and r[time_idx] is not None]
            if not samples or pd.to_datetime(samples, errors="coerce").isna().all():
                errors.append("Time column has no parseable timestamp in the sampled rows")
        summary = {
            "cars": sorted(cars, key=lambda car: (int(car), car)),
            "recognized_fields": recognized,
            "sampled_rows": max(0, len(rows) - 1),
        }
        return {"valid": not errors, "errors": errors, "summary": summary}
    finally:
        book.close()


@router.post("/acv/validate")
async def validate_acv_upload(file: UploadFile = File(...)) -> dict[str, Any]:
    """Lightweight preflight for the ACV workbook picker, without creating a result session."""
    name = Path(str(file.filename or "")).name
    if not name.lower().endswith(".xlsx") or name.startswith("~$"):
        await file.close()
        return {"valid": False, "errors": ["choose a .xlsx ACV case workbook"], "summary": {}}
    root = _upload_root()
    root.mkdir(parents=True, exist_ok=True)
    with tempfile.TemporaryDirectory(prefix="acv-check-", dir=root) as directory:
        dest = Path(directory) / name
        await _save_upload(file, dest)
        return await run_in_threadpool(inspect_acv_workbook, dest)


@router.get("/tasks")
def ps3_tasks(request: Request) -> list[dict[str, Any]]:
    """The four subsystems, what they accept, and whether each can actually run right now."""
    _sweep()
    local_paths = (os.environ.get("NEBULAX_LOCAL_PATHS") == "1"
                   and request.client is not None and request.client.host in ("127.0.0.1", "::1")
                   and request.url.hostname in ("127.0.0.1", "localhost", "::1"))
    out: list[dict[str, Any]] = []
    for name in TASK_NAMES:
        try:
            resolve_task(name)
        except TaskUnavailable as exc:
            available, detail = False, str(exc)
        else:
            available, detail = True, None
        out.append(
            {
                "name": name,
                "label": TASK_LABELS[name],
                "accepts": list(ACCEPTED_SUFFIXES[name]),
                "output_filename": OUTPUT_FILENAMES[name],
                "cv": cv_summary(name),
                "model_loaded": common.model_path(name, model_dir=_model_dir()).exists(),
                "max_file_bytes": MAX_UPLOAD_BYTES,
                "max_files_per_request": MAX_BATCH_FILES,
                "local_paths": local_paths,
                "available": available,
                "detail": detail,
            }
        )
    return out


@router.get("/results/{session}.csv")
def ps3_results_csv(session: str) -> Response:
    """The accumulated session as the organiser CSV - byte-identical to the CLI's output."""
    _sweep()
    sess = _session_or_404(session)
    sess.touched = time.time()
    if not sess.rows:
        raise HTTPException(status_code=404, detail=f"session {session!r} has no predictions yet")
    text = render_csv(sess.task, sess.rows)
    tmp = sess.dir / OUTPUT_FILENAMES[sess.task]
    tmp.parent.mkdir(parents=True, exist_ok=True)
    tmp.write_text(text, encoding="utf-8", newline="")
    report = validate_csv(
        sess.task,
        tmp,
        None,
        expected_cars=sess.expected_cars or None,
    )
    if not report.ok:
        raise HTTPException(
            status_code=500,
            detail=f"the accumulated {sess.task} predictions are not a valid submission: " + "; ".join(report.errors),
        )
    return Response(
        content=text.encode("utf-8"),
        media_type="text/csv; charset=utf-8",
        headers={
            "Content-Disposition": f'attachment; filename="{OUTPUT_FILENAMES[sess.task]}"',
            "X-PS3-Rows": str(len(sess.rows)),
            "X-PS3-Warnings": str(len(report.warnings)),
        },
    )


@router.delete("/results/{session}")
def ps3_results_delete(session: str) -> dict[str, Any]:
    """Forget a session and delete every temp file it uploaded."""
    _sweep()
    sess = _session_or_404(session)
    shutil.rmtree(sess.dir, ignore_errors=True)
    _SESSIONS.pop(sess.token, None)
    return {"session": sess.token, "deleted": True, "n_files": len(sess.files), "n_rows": len(sess.rows)}


@router.get("/{task}/cv")
def ps3_cv(task: str) -> dict[str, Any]:
    """``results/ps3/<task>_cv.json`` as produced by ``scripts/ps3_train.py``."""
    key = _task_or_404(task)
    path = cv_path(key)
    if not path.exists():
        raise HTTPException(
            status_code=404,
            detail=f"no CV results for {key!r} yet ({path}); run `python scripts/ps3_train.py --task {key}`",
        )
    try:
        return json.loads(path.read_text(encoding="utf-8"))
    except (OSError, ValueError) as exc:
        raise HTTPException(status_code=500, detail=f"{path} is not readable json: {exc}") from exc


async def _extract_uploads(
    request: Request, files: list[UploadFile] | None, session: str | None
) -> tuple[list[UploadFile], str | None]:
    uploads = list(files or [])
    if not uploads:  # tolerate the HTML `files[]` spelling (and a single `file`)
        form = await request.form()
        uploads = [v for k, v in form.multi_items() if k in ("files[]", "file") and isinstance(v, _StarletteUpload)]
        if session is None:
            raw = form.get("session")
            session = raw if isinstance(raw, str) else None
    if session in (None, ""):
        session = request.query_params.get("session") or None
    return uploads, session


async def _run_predict_batch(
    key: str,
    uploads: list[UploadFile],
    session: str | None,
) -> tuple[_Session, list[dict[str, Any]], list[dict[str, str]], list[Path], Any]:
    try:
        task_obj = resolve_task(key)
    except TaskUnavailable as exc:
        raise HTTPException(status_code=503, detail=str(exc)) from exc

    if session:
        sess = _session_or_404(session)
        if sess.task != key:
            raise HTTPException(
                status_code=400,
                detail=f"session {session!r} belongs to task {sess.task!r}, not {key!r}",
            )
    else:
        now = time.time()
        token = uuid.uuid4().hex
        sess = _Session(token=token, task=key, created=now, touched=now)
        _SESSIONS[token] = sess
    sess.touched = time.time()
    sess.n_batches += 1
    batch_dir = sess.dir / f"batch{sess.n_batches:03d}"
    batch_dir.mkdir(parents=True, exist_ok=True)

    suffixes = ACCEPTED_SUFFIXES[key]
    errors: list[dict[str, str]] = []
    explanations: list[dict[str, Any]] = []
    saved_paths: list[Path] = []
    model = task_model(key)
    n_ok = 0
    for upload in uploads:
        name = Path(str(upload.filename or "")).name.strip()
        if not name or name in (".", ".."):
            errors.append({"file": str(upload.filename or ""), "message": "upload has no usable file name"})
            continue
        if Path(name).suffix.lower() not in suffixes:
            errors.append(
                {
                    "file": name,
                    "message": f"{TASK_LABELS[key]} accepts {' or '.join(suffixes)}, not {Path(name).suffix or 'a file with no extension'}",
                }
            )
            continue
        if name in sess.files:
            errors.append({"file": name, "message": "already predicted in this session; delete the session to start over"})
            continue
        dest = batch_dir / name
        await _save_upload(upload, dest)
        try:
            result = await run_in_threadpool(run_file, task_obj, dest, model)
            rows = task_obj.to_rows(result)
            explanation = task_obj.explain(result).as_dict()
        except FileNotFoundError as exc:
            raise HTTPException(
                status_code=503,
                detail=f"PS3 task {key!r} cannot predict yet: {exc}",
            ) from exc
        except Exception as exc:
            errors.append({"file": name, "message": f"{type(exc).__name__}: {exc}"})
            continue
        sess.rows.extend(rows)
        sess.files.append(name)
        explanations.append(explanation)
        saved_paths.append(dest)
        if key == "acv":
            try:
                sess.expected_cars[name] = acv_car_ids(dest)
            except Exception:  # a car-id sniff must never fail a prediction
                pass
        n_ok += 1

    if n_ok == 0:
        raise HTTPException(
            status_code=400,
            detail="no file in this batch could be predicted: "
            + "; ".join(f"{e['file']}: {e['message']}" for e in errors),
        )
    return sess, explanations, errors, saved_paths, model


def _local_paths(request: Request, task: str, raw_path: str) -> list[Path]:
    """Resolve an explicitly entered path for the loopback-only desktop app."""
    if (os.environ.get("NEBULAX_LOCAL_PATHS") != "1"
        or request.client is None or request.client.host not in ("127.0.0.1", "::1")
        or request.url.hostname not in ("127.0.0.1", "localhost", "::1")):
        raise HTTPException(status_code=403, detail="local path import is only available in the desktop app")
    origin = request.headers.get("origin")
    if origin and origin != f"{request.url.scheme}://{request.headers.get('host', '')}":
        raise HTTPException(status_code=403, detail="local path import requires the app's own page")
    entered = str(raw_path or "").strip()
    if not entered or not (entered.startswith("/") or entered.startswith("~/")):
        raise HTTPException(status_code=400, detail="enter an absolute file or folder path, or one starting with ~/")
    target = Path(entered).expanduser().resolve()
    try:
        paths = collect_inputs(task, target) if target.is_dir() else [target]
    except FileNotFoundError as exc:
        raise HTTPException(status_code=400, detail=str(exc)) from exc
    if not paths or len(paths) > 500:
        raise HTTPException(status_code=400, detail="choose a folder with 1 to 500 matching files")
    if task == "door" and len(paths) != 1:
        raise HTTPException(status_code=400, detail="Door accepts one Test.csv stream at a time")
    for path in paths:
        if not path.is_file() or path.suffix.lower() not in ACCEPTED_SUFFIXES[task]:
            raise HTTPException(status_code=400, detail=f"{path.name}: not a readable {TASK_LABELS[task]} input")
        try:
            size = path.stat().st_size
            with path.open("rb"):
                pass
        except OSError as exc:
            raise HTTPException(status_code=400, detail=f"cannot read {path.name}: {exc}") from exc
        if size > MAX_UPLOAD_BYTES:
            raise HTTPException(status_code=413, detail=f"{path.name} is larger than the {MAX_UPLOAD_BYTES // (1024 * 1024)} MB per-file limit")
        if path.suffix.lower() == ".csv":
            try:
                detected = infer_task(path)
            except (OSError, ValueError) as exc:
                raise HTTPException(status_code=400, detail=f"{path.name}: cannot identify the CSV system: {exc}") from exc
            if detected != task:
                raise HTTPException(
                    status_code=400,
                    detail=f"{path.name} looks like {TASK_LABELS[detected]} data. Select the {TASK_LABELS[detected]} tab, then ADD PATH.",
                )
    return paths


@router.post("/{task}/local/inspect")
def ps3_local_inspect(task: str, request: Request, body: dict[str, Any] = Body(...)) -> dict[str, Any]:
    """List local files for Firefox when its Snap file chooser does not return selections."""
    key = _task_or_404(task)
    paths = _local_paths(request, key, body.get("path", ""))
    if key == "acv":
        for path in paths:
            report = inspect_acv_workbook(path)
            if not report["valid"]:
                raise HTTPException(status_code=400, detail=f"{path.name}: {'; '.join(report['errors'])}")
    return {"files": [{"name": p.name, "path": str(p), "size": p.stat().st_size} for p in paths]}


@router.post("/{task}/local/stream")
async def ps3_local_stream(task: str, request: Request, body: dict[str, Any] = Body(...)) -> dict[str, Any]:
    key = _task_or_404(task)
    paths = _local_paths(request, key, body.get("path", ""))
    if len(paths) != 1:
        raise HTTPException(status_code=400, detail="stream one local file at a time")
    path = paths[0]
    with path.open("rb") as fh:
        upload = UploadFile(file=fh, filename=path.name)
        return await ps3_stream(key, request, files=[upload], session=body.get("session"))


@router.post("/{task}/predict")
async def ps3_predict(
    task: str,
    request: Request,
    files: list[UploadFile] | None = File(default=None),
    session: str | None = Form(default=None),
) -> dict[str, Any]:
    """Predict one **batch** of uploads, appending to a session.

    The client repeats the call with the returned ``session`` for the next batch, so the 68 rail
    files (~1.1 GB) never travel in one request - at most ``MAX_BATCH_FILES`` per call, each at
    most ``MAX_UPLOAD_BYTES``. Files are processed one at a time, in the order they were sent.
    """
    _sweep()
    key = _task_or_404(task)
    uploads, session = await _extract_uploads(request, files, session)
    if not uploads:
        raise HTTPException(status_code=400, detail="no files uploaded (send them as multipart `files`)")
    if len(uploads) > MAX_BATCH_FILES:
        raise HTTPException(
            status_code=413,
            detail=(
                f"{len(uploads)} files in one request; send at most {MAX_BATCH_FILES} per request "
                "and pass the returned `session` back with the next batch"
            ),
        )

    sess, explanations, errors, _, _ = await _run_predict_batch(key, uploads, session)
    return {
        "session": sess.token,
        "task": key,
        "output_filename": OUTPUT_FILENAMES[key],
        "rows": [dict(r) for r in sess.rows],
        "explanations": explanations,
        "csv_url": f"/api/ps3/results/{sess.token}.csv",
        "n_done": len(sess.files),
        "files": list(sess.files),
        "errors": errors,
        "expires_at": sess.touched + SESSION_TTL_S,
    }


@router.post("/{task}/stream")
async def ps3_stream(
    task: str,
    request: Request,
    files: list[UploadFile] | None = File(default=None),
    session: str | None = Form(default=None),
) -> dict[str, Any]:
    """Predict and stream exactly one file, returning preview frames plus the session rows."""
    from nebulax.ps3.stream import stream_file

    _sweep()
    key = _task_or_404(task)
    uploads, session = await _extract_uploads(request, files, session)
    if len(uploads) != 1:
        raise HTTPException(
            status_code=400,
            detail=f"stream expects exactly one file, got {len(uploads)}; send as multipart `files`",
        )

    previous = _session_or_404(session) if session else None
    snapshot = (
        (list(previous.rows), list(previous.files), dict(previous.expected_cars), previous.n_batches, previous.touched)
        if previous else None
    )
    old_tokens = set(_SESSIONS)
    try:
        sess, explanations, errors, saved_paths, model = await _run_predict_batch(key, uploads, session)
        # a refused upload (wrong suffix, or a name already predicted in this session after the
        # page's STOP cut the reply off) has no saved path: report it in `errors`, with no frames
        frames = (
            await run_in_threadpool(stream_file, key, saved_paths[0], model, max_frames=60)
            if saved_paths else []
        )
    except Exception:
        if previous and snapshot:
            batch_number = previous.n_batches
            previous.rows, previous.files, previous.expected_cars, previous.n_batches, previous.touched = snapshot
            shutil.rmtree(previous.dir / f"batch{batch_number:03d}", ignore_errors=True)
        else:
            for token in set(_SESSIONS) - old_tokens:
                shutil.rmtree(_SESSIONS[token].dir, ignore_errors=True)
                _SESSIONS.pop(token, None)
        raise
    return {
        "session": sess.token,
        "task": key,
        "output_filename": OUTPUT_FILENAMES[key],
        "rows": [dict(r) for r in sess.rows],
        "explanations": explanations,
        "csv_url": f"/api/ps3/results/{sess.token}.csv",
        "n_done": len(sess.files),
        "files": list(sess.files),
        "errors": errors,
        "frames": [f.as_dict() for f in frames],
        "expires_at": sess.touched + SESSION_TTL_S,
    }


def _example_stream(key: str) -> dict[str, Any]:
    """Select one deterministic predicted example from the local organiser Test data."""
    from nebulax.ps3.stream import stream_file

    try:
        task = resolve_task(key)
        model = task_model(key)
    except TaskUnavailable as exc:
        raise HTTPException(status_code=503, detail=str(exc)) from exc
    root = common.test_dir(key)
    try:
        paths = [root / "Test.csv"] if key == "door" else collect_inputs(key, root)
    except FileNotFoundError as exc:
        raise HTTPException(status_code=404, detail=f"local {key} Test data is unavailable") from exc
    if not paths or not paths[0].is_file():
        raise HTTPException(status_code=404, detail=f"local {key} Test data is unavailable")

    chosen: tuple[Path, PredictionResult, list[dict[str, Any]]] | None = None
    best_damage = -math.inf
    try:
        for path in paths:
            result = run_file(task, path, model)
            rows = task.to_rows(result)
            if key == "door":
                if any("abnormal" in str(r.get("prediction", "")).lower() for r in rows):
                    chosen = (path, result, rows)
                break
            if key == "acv":
                if rows and str(rows[0].get("ranked_cars", "")).strip():
                    chosen = (path, result, rows)
                break
            if key == "rail":
                if rows and str(rows[0].get("prediction", "")).strip().lower() in ("side i", "side ii"):
                    chosen = (path, result, rows)
                    break
            if key == "shm" and rows:
                damage = float(rows[0].get("prediction", "nan"))
                if math.isfinite(damage) and damage > best_damage:
                    best_damage = damage
                    chosen = (path, result, rows)
        if chosen is None:
            raise HTTPException(status_code=409, detail=f"no qualifying predicted {key} example in local Test data")
        path, result, rows = chosen
        frames = stream_file(key, path, model, max_frames=60)
    except HTTPException:
        raise
    except FileNotFoundError as exc:
        raise HTTPException(status_code=503, detail=f"PS3 task {key!r} cannot predict yet: {exc}") from exc
    except Exception as exc:
        raise HTTPException(status_code=503, detail=f"could not prepare {key} example: {type(exc).__name__}: {exc}") from exc

    if key == "door":
        abnormal = next(r for r in rows if "abnormal" in str(r.get("prediction", "")).lower())
        frame_index = next((i for i, f in enumerate(frames) if abnormal in f.rows), len(frames) - 1)
        criterion = "first predicted abnormal-resistance cycle"
    elif key == "acv":
        frame_index = len(frames) - 1
        criterion = f"final predicted rank-one car {str(rows[0]['ranked_cars']).split('|')[0]}"
    elif key == "rail":
        frame_index = len(frames) - 1
        criterion = "first Test file predicted Side I or Side II"
    else:
        frame_index = len(frames) - 1
        criterion = "highest final predicted fatigue damage"
    return {
        "task": key,
        "source": "local organiser Test data",
        "files": [path.name],
        "rows": [dict(r) for r in rows],
        "explanations": [task.explain(result).as_dict()],
        "frames": [f.as_dict() for f in frames],
        "example": {"file": path.name, "frame_index": frame_index, "criterion": criterion, "prediction_only": True},
    }


@router.get("/{task}/example/stream")
async def ps3_example_stream(task: str) -> dict[str, Any]:
    """Replay a predicted example; no labels, arbitrary path, or upload session are exposed."""
    key = _task_or_404(task)
    return await run_in_threadpool(_example_stream, key)
