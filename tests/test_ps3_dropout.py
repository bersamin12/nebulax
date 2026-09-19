"""Simulated missing columns (nebulax.ps3.dropout) and the `drop_columns` upload field."""

from __future__ import annotations

import csv
import shutil
from pathlib import Path

import pytest
from openpyxl import load_workbook

from nebulax.ps3 import dropout
from nebulax.ps3.acv_features import load_case
from nebulax.ps3.door_features import load_stream

# the router fixtures (isolated upload dirs, a TestClient, DummyTask registration)
from tests.test_api_ps3 import _isolated, client, task_factory  # noqa: F401

FIXTURES = Path(__file__).resolve().parent / "fixtures" / "ps3"
DOOR_CSV = FIXTURES / "door" / "train_slice.csv"
ACV_XLSX = FIXTURES / "acv" / "case01_slice.xlsx"


def _header(path: Path) -> list[str]:
    with open(path, encoding="utf-8-sig", newline="") as fh:
        return next(csv.reader(fh))


def test_droppable_fields_exclude_the_required_ones() -> None:
    door = dropout.droppable_fields("door")
    assert "voltage" in door and "emf" in door
    assert not {"datetime", "current", "position"} & set(door)
    acv = dropout.droppable_fields("acv")
    assert "outdoor" in acv and "indoor" not in acv
    assert dropout.droppable_fields("rail") == ()
    assert dropout.droppable_fields("shm") == ()


def test_parse_drop_list_normalises_and_rejects_unknown_fields() -> None:
    assert dropout.parse_drop_list("door", None) == []
    assert dropout.parse_drop_list("door", " Voltage, emf ,voltage") == ["voltage", "emf"]
    with pytest.raises(ValueError, match="cannot drop current"):
        dropout.parse_drop_list("door", "current")
    with pytest.raises(ValueError, match="no optional columns"):
        dropout.parse_drop_list("rail", "speed")


def test_door_columns_are_removed_and_the_stream_still_loads(tmp_path: Path) -> None:
    p = tmp_path / "Test.csv"
    shutil.copy(DOOR_CSV, p)
    before = _header(p)
    removed = dropout.apply_column_dropout("door", p, ["voltage", "emf"])
    assert removed == ["Motor Voltage(10mV)", "Motor electrodynamic force"]
    after = _header(p)
    assert len(after) == len(before) - 2 and not set(removed) & set(after)
    # the loader accepts the thinner file; the dropped fields come back as NaN
    stream = load_stream(p)
    assert stream.table["voltage"].isna().all() and stream.table["emf"].isna().all()
    assert stream.table["current"].notna().any()
    # a second pass finds nothing left to remove and leaves the file alone
    assert dropout.apply_column_dropout("door", p, ["voltage"]) == []


def test_acv_columns_are_removed_for_every_car(tmp_path: Path) -> None:
    p = tmp_path / "case.xlsx"
    shutil.copy(ACV_XLSX, p)
    removed = dropout.apply_column_dropout("acv", p, ["outdoor"])
    assert len(removed) == 8 and all("Outdoor" in name or "Outside" in name or "Fresh Air" in name for name in removed)
    header = [c.value for c in next(load_workbook(p, read_only=True).worksheets[0].iter_rows(max_row=1))]
    assert not any(name in header for name in removed)
    case = load_case(p)  # the loader falls back to the indoor median for the hot rows
    assert case.panel["indoor"].notna().any().any()
    assert "outdoor" not in case.panel or case.panel["outdoor"].isna().all().all()


def test_api_predict_accepts_drop_columns(client, task_factory) -> None:
    task_factory("door")
    body = client.post(
        "/api/ps3/door/predict",
        files=[("files", (DOOR_CSV.name, DOOR_CSV.read_bytes(), "application/octet-stream"))],
        data={"drop_columns": "voltage,emf"},
    ).json()
    assert body["explanations"][0]["dropped_fields"] == ["voltage", "emf"]
    assert body["explanations"][0]["dropped_columns"] == ["Motor Voltage(10mV)", "Motor electrodynamic force"]


def test_api_predict_rejects_a_required_or_unknown_field(client, task_factory) -> None:
    task_factory("door")
    r = client.post(
        "/api/ps3/door/predict",
        files=[("files", (DOOR_CSV.name, DOOR_CSV.read_bytes(), "application/octet-stream"))],
        data={"drop_columns": "current"},
    )
    assert r.status_code == 400 and "cannot drop current" in r.json()["detail"]


def test_api_tasks_advertise_droppable_fields(client, task_factory) -> None:
    task_factory("door")
    task_factory("rail")
    tasks = {t["name"]: t for t in client.get("/api/ps3/tasks").json()}
    assert "voltage" in tasks["door"]["droppable_fields"]
    assert tasks["rail"]["droppable_fields"] == []
