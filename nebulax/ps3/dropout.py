"""Simulate missing columns in an uploaded Door stream or ACV workbook.

The upload page lets a user *deactivate* optional fields before RUN so they can see how a model
copes with less data than the released files carry. The simulation is done on the saved input
itself, before the task's normal ``load``: the columns are physically removed from the file, so
the loaders' own fallbacks (Door: NaN features filled with their neutral value; ACV: the
outdoor -> indoor-median -> all-rows hot-row fallback) take over exactly as they would for a
fleet that never logged the field.

Only Door and ACV are droppable. Rail needs all 129 columns (``rail_features.read_rail_csv``
rejects anything else, and the same-side coherence needs every box) and SHM is one column.
"""

from __future__ import annotations

import csv
from pathlib import Path
from typing import Iterable

__all__ = ["DROPPABLE", "droppable_fields", "parse_drop_list", "apply_column_dropout"]


def _door_droppable() -> tuple[str, ...]:
    from nebulax.ps3.door_features import _REQUIRED, CANONICAL_COLUMNS

    return tuple(c for c in CANONICAL_COLUMNS if c not in _REQUIRED)


def _acv_droppable() -> tuple[str, ...]:
    from nebulax.ps3.acv_features import SYNONYMS

    # ``indoor`` is the one field the ranking cannot do without
    return tuple(c for c in SYNONYMS if c != "indoor")


#: Canonical fields a user may switch off, per task (lazy: the feature modules are heavy).
DROPPABLE = {"door": _door_droppable, "acv": _acv_droppable}


def droppable_fields(task: str) -> tuple[str, ...]:
    """The canonical fields that may be dropped for ``task`` (empty for rail / shm)."""
    fn = DROPPABLE.get(task)
    return fn() if fn else ()


def parse_drop_list(task: str, raw: str | Iterable[str] | None) -> list[str]:
    """Normalise a comma-separated (or iterable) drop request; unknown fields raise ``ValueError``."""
    if raw is None:
        return []
    items = raw.split(",") if isinstance(raw, str) else list(raw)
    wanted = [s.strip().lower() for s in items if s and s.strip()]
    if not wanted:
        return []
    allowed = droppable_fields(task)
    if not allowed:
        raise ValueError(f"{task!r} has no optional columns to drop")
    bad = sorted(set(wanted) - set(allowed))
    if bad:
        raise ValueError(f"cannot drop {', '.join(bad)} for {task!r}; optional fields are {', '.join(allowed)}")
    return list(dict.fromkeys(wanted))  # de-duplicated, order kept


def _drop_door_columns(path: Path, fields: list[str]) -> list[str]:
    from nebulax.ps3.door_features import CANONICAL_COLUMNS, _norm_key

    aliases = {alias for f in fields for alias in CANONICAL_COLUMNS[f]}
    tmp = path.with_suffix(path.suffix + ".drop")
    with open(path, "r", encoding="utf-8-sig", newline="") as src:
        reader = csv.reader(src)
        header = next(reader, None)
        if header is None:
            return []
        keep = [i for i, name in enumerate(header) if _norm_key(name) not in aliases]
        removed = [name for i, name in enumerate(header) if i not in keep]
        if not removed:
            return []
        with open(tmp, "w", encoding="utf-8", newline="") as out:
            writer = csv.writer(out)
            writer.writerow([header[i] for i in keep])
            for row in reader:
                writer.writerow([row[i] if i < len(row) else "" for i in keep])
    tmp.replace(path)
    return removed


def _drop_acv_columns(path: Path, fields: list[str]) -> list[str]:
    from openpyxl import load_workbook

    from nebulax.ps3.acv_features import CAR_COL_RE, SYNONYMS

    params = {syn.lower() for f in fields for syn in SYNONYMS[f]}
    book = load_workbook(path)
    try:
        sheet = book.worksheets[0]
        removed: list[str] = []
        hits: list[int] = []
        for cell in next(sheet.iter_rows(min_row=1, max_row=1)):
            name = str(cell.value).strip() if cell.value is not None else ""
            m = CAR_COL_RE.match(name)
            if m and m.group(2).strip().lower() in params:
                hits.append(cell.column)
                removed.append(name)
        for col in sorted(hits, reverse=True):  # right to left so the indices stay valid
            sheet.delete_cols(col)
        if removed:
            book.save(path)
    finally:
        book.close()
    return removed


def apply_column_dropout(task: str, path: Path | str, fields: Iterable[str]) -> list[str]:
    """Remove the columns behind ``fields`` from the file at ``path``, in place.

    Returns the raw column names that were removed (empty when the file did not carry them).
    ``fields`` are canonical names from :func:`droppable_fields`; anything else raises
    ``ValueError`` before the file is touched.
    """
    wanted = parse_drop_list(task, list(fields))
    if not wanted:
        return []
    p = Path(path)
    if task == "door":
        return _drop_door_columns(p, wanted)
    if task == "acv":
        return _drop_acv_columns(p, wanted)
    raise ValueError(f"{task!r} has no optional columns to drop")
